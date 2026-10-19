#!/usr/bin/env python3
"""Enrich enterprise sample CSVs through the webSearchAgent API.

The original CSV/XLSX files are read-only inputs. Network access only happens in
normal mode; --dry-run and --self-test never call the API or create output files.
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import os
import re
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import warnings
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable


FILE_STEM = "ocr_review_명함정리_샘플"
SAMPLE_SHEETS = tuple(f"샘플{number}" for number in range(1, 6))

TARGET_COLUMNS = {
    "company": "식별정보(회사명)",
    "business_no": "사업자등록번호",
    "address": "식별정보(주소)",
    "representative": "식별정보(대표자명)",
    "products": "식별정보(주요상품)",
    "homepage": "식별정보(홈페이지)",
    "scale": "식별정보(기업규모)",
    "type": "식별정보(기업유형)",
    "industry": "식별정보(업종)",
    "founded": "식별정보(설립일)",
    "revenue": "식별정보(매출액)",
    "operating_profit": "식별정보(영업이익)",
    "net_income": "식별정보(당기순이익)",
}

API_OUTPUT_NAMES = {
    "business_no": "사업자번호",
}

GENERATED_COLUMNS = {TARGET_COLUMNS["business_no"]}

FIELD_ALIASES = {
    "company": ("기업명", "회사명", "법인명", "상호명"),
    "business_no": ("사업자등록번호", "사업자번호", "사업자등록 No", "Business Registration Number"),
    "address": ("본사주소", "기업주소", "사업장주소", "소재지", "주소"),
    "representative": ("대표자명", "대표자", "CEO"),
    "products": ("주요사업", "주요제품", "주요상품", "사업내용", "주요서비스"),
    "homepage": ("홈페이지", "웹사이트", "Website", "사이트"),
    "scale": ("기업규모", "사원수", "직원수", "종업원수"),
    "type": ("기업형태", "기업유형", "법인구분"),
    "industry": ("산업", "업종", "업태", "종목"),
    "founded": ("설립일자", "설립일", "설립연도", "설립"),
    "revenue": ("연매출", "매출액", "매출"),
    "operating_profit": ("영업이익",),
    "net_income": ("당기순이익", "순이익"),
}

META_COLUMNS = (
    "보강상태",
    "정합성검토",
    "보강필드",
    "보강출처",
    "보강근거",
    "보강검색어",
    "보강시각(UTC)",
    "robots.txt 준수",
    "보강오류",
)

KNOWN_ENTERPRISE_DOMAINS = (
    "jobkorea.co.kr",
    "saramin.co.kr",
    "catch.co.kr",
    "incruit.com",
    "wanted.co.kr",
    "rocketpunch.com",
    "nicebizinfo.com",
    "bizno.net",
    "bizline.co.kr",
    "biznet.co.kr",
    "kreport.co.kr",
    "dart.fss.or.kr",
    "kind.krx.co.kr",
)

IGNORED_URL_MARKERS = (
    "/login",
    "/signin",
    "/captcha",
    "/search?",
    "/search/",
    "accounts.google.",
)

VALUE_LIMITS = {
    "company": 120,
    "business_no": 20,
    "address": 220,
    "representative": 80,
    "products": 300,
    "homepage": 220,
    "scale": 100,
    "type": 100,
    "industry": 160,
    "founded": 80,
    "revenue": 120,
    "operating_profit": 120,
    "net_income": 120,
}

EMPTY_VALUES = {
    "",
    "-",
    "--",
    "없음",
    "미상",
    "비공개",
    "정보없음",
    "정보 없음",
    "해당없음",
    "n/a",
    "null",
    "none",
}

PUBLIC_EMAIL_DOMAINS = {
    "daum.net",
    "gmail.com",
    "googlemail.com",
    "hanmail.net",
    "hotmail.com",
    "icloud.com",
    "kakao.com",
    "korea.com",
    "live.com",
    "msn.com",
    "nate.com",
    "naver.com",
    "outlook.com",
    "yahoo.com",
}


class ApiError(RuntimeError):
    pass


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def digits(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def format_business_number(value: Any) -> str:
    value = digits(value)
    return f"{value[:3]}-{value[3:5]}-{value[5:]}" if len(value) == 10 else ""


def business_numbers(value: Any) -> set[str]:
    return {
        "".join(match)
        for match in re.findall(r"(?<!\d)(\d{3})[-\s]?(\d{2})[-\s]?(\d{5})(?!\d)", str(value or ""))
    }


def normalize_text(value: Any) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]", "", str(value or "")).lower()


def normalize_company(value: Any) -> str:
    normalized = normalize_text(value)
    for legal_term in (
        "주식회사",
        "유한회사",
        "유한책임회사",
        "사단법인",
        "재단법인",
        "corporation",
        "companylimited",
        "coltd",
        "incorporated",
    ):
        normalized = normalized.replace(legal_term, "")
    return normalized


def company_names_equivalent(left: str, right: str) -> bool:
    """Strict identity check for extracted legal names, not loose SERP ranking."""
    def identity_name(value: str) -> str:
        without_short_legal_name = re.sub(
            r"\(\s*주\s*\)|㈜", "", str(value or ""), flags=re.IGNORECASE
        )
        return normalize_company(without_short_legal_name)

    left_normalized = identity_name(left)
    right_normalized = identity_name(right)
    return bool(left_normalized and left_normalized == right_normalized)


def corporate_email_host(value: Any) -> str:
    match = re.search(r"[A-Z0-9._%+-]+@([A-Z0-9.-]+\.[A-Z]{2,})", str(value or ""), re.IGNORECASE)
    host = (match.group(1).lower().removeprefix("www.") if match else "")
    return "" if host in PUBLIC_EMAIL_DOMAINS else host


def parse_bool(value: str | bool) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "on"}:
        return True
    if normalized in {"0", "false", "no", "off"}:
        return False
    raise argparse.ArgumentTypeError("expected true or false")


def existing_file(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise argparse.ArgumentTypeError(f"file not found: {path}")
    return path


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Enrich 샘플1~샘플5 and create a separate PM-delivery XLSX.",
    )
    default_dir = Path(__file__).resolve().parent
    parser.add_argument("--input-xlsx", type=existing_file, default=default_dir / f"{FILE_STEM}.xlsx")
    parser.add_argument("--csv-dir", type=Path, default=default_dir)
    parser.add_argument("--output-xlsx", type=Path, default=default_dir / f"{FILE_STEM}_websearch_enriched.xlsx")
    parser.add_argument("--progress-file", type=Path, default=default_dir / "enterprise_enrichment_progress.jsonl")
    parser.add_argument("--base-url", default="http://127.0.0.1:9101")
    parser.add_argument("--provider", default="serper")
    parser.add_argument("--crawler", default="auto")
    parser.add_argument("--respect-robots", type=parse_bool, default=True)
    parser.add_argument("--search-limit", type=int, default=5)
    parser.add_argument("--crawl-limit", type=int, default=2)
    parser.add_argument("--host-delay-seconds", type=float, default=2.0)
    parser.add_argument("--request-timeout", type=float, default=240.0)
    parser.add_argument("--workers", type=int, default=1)
    parser.add_argument("--max-companies", type=int, default=0, help="0 means all companies")
    parser.add_argument("--refresh", action="store_true", help="Ignore the JSONL resume cache")
    parser.add_argument("--overwrite", action="store_true", help="Replace an existing generated XLSX")
    parser.add_argument("--dry-run", action="store_true", help="Validate and print the plan; no network or writes")
    parser.add_argument("--self-test", action="store_true", help=argparse.SUPPRESS)
    return parser


def validate_args(args: argparse.Namespace, parser: argparse.ArgumentParser) -> None:
    args.csv_dir = args.csv_dir.expanduser().resolve()
    args.output_xlsx = args.output_xlsx.expanduser().resolve()
    args.progress_file = args.progress_file.expanduser().resolve()
    args.base_url = args.base_url.rstrip("/")
    parsed_base = urllib.parse.urlsplit(args.base_url)
    if parsed_base.scheme not in {"http", "https"} or not parsed_base.netloc:
        parser.error("--base-url must be an http(s) URL")
    if args.input_xlsx.resolve() == args.output_xlsx:
        parser.error("output XLSX must differ from the original input XLSX")
    if not 1 <= args.search_limit <= 20:
        parser.error("--search-limit must be between 1 and 20")
    if not 1 <= args.crawl_limit <= args.search_limit + 1:
        parser.error("--crawl-limit must be between 1 and search-limit + 1")
    if not 0 <= args.host_delay_seconds <= 60:
        parser.error("--host-delay-seconds must be between 0 and 60")
    if args.request_timeout <= 0:
        parser.error("--request-timeout must be positive")
    if not 1 <= args.workers <= 4:
        parser.error("--workers must be between 1 and 4")
    if args.max_companies < 0:
        parser.error("--max-companies cannot be negative")
    if args.output_xlsx.exists() and not (args.overwrite or args.dry_run):
        parser.error(f"output already exists; pass --overwrite to replace it: {args.output_xlsx}")


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.is_file():
        raise FileNotFoundError(f"required CSV not found: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = list(reader.fieldnames or [])
        rows = [{key: "" if value is None else value for key, value in row.items()} for row in reader]
    if not headers or not rows:
        raise ValueError(f"CSV is empty: {path}")
    return headers, rows


def load_inputs(csv_dir: Path) -> tuple[list[str], dict[str, list[dict[str, str]]]]:
    original_csv = csv_dir / f"{FILE_STEM}_원본.csv"
    read_csv(original_csv)

    common_headers: list[str] | None = None
    samples: dict[str, list[dict[str, str]]] = {}
    required = {
        "회사명(Company)",
        "회사명_영문(Alt Company)",
        "식별여부",
        *(column for column in TARGET_COLUMNS.values() if column not in GENERATED_COLUMNS),
    }
    for sheet in SAMPLE_SHEETS:
        headers, rows = read_csv(csv_dir / f"{FILE_STEM}_{sheet}.csv")
        missing = sorted(required - set(headers))
        if missing:
            raise ValueError(f"{sheet} is missing columns: {', '.join(missing)}")
        if common_headers is None:
            common_headers = headers
        elif headers != common_headers:
            raise ValueError(f"column order differs in {sheet}")
        samples[sheet] = rows
    return common_headers or [], samples


def validate_workbook(path: Path, headers: list[str], samples: dict[str, list[dict[str, str]]]) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required: python3 -m pip install openpyxl") from exc

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        required_sheets = {"원본", *SAMPLE_SHEETS}
        missing = sorted(required_sheets - set(workbook.sheetnames))
        if missing:
            raise ValueError(f"input workbook is missing sheets: {', '.join(missing)}")
        for sheet in SAMPLE_SHEETS:
            worksheet = workbook[sheet]
            actual_headers = [clean(worksheet.cell(1, column).value) for column in range(1, len(headers) + 1)]
            if actual_headers != headers:
                raise ValueError(f"XLSX/CSV header mismatch in {sheet}")
            if worksheet.max_row != len(samples[sheet]) + 1:
                raise ValueError(
                    f"XLSX/CSV row-count mismatch in {sheet}: "
                    f"xlsx={worksheet.max_row - 1}, csv={len(samples[sheet])}"
                )
    finally:
        workbook.close()


def company_key(row: dict[str, str]) -> str:
    business_no = digits(row.get("식별여부"))
    if len(business_no) == 10:
        return f"biz:{business_no}"
    company = row.get(TARGET_COLUMNS["company"]) or row.get("회사명(Company)")
    normalized = normalize_company(company)
    if not normalized:
        normalized = normalize_text(row.get("이미지") or json.dumps(row, ensure_ascii=False, sort_keys=True))
    return f"name:{normalized}"


def group_entities(samples: dict[str, list[dict[str, str]]]) -> dict[str, list[dict[str, str]]]:
    entities: dict[str, list[dict[str, str]]] = defaultdict(list)
    for sheet, rows in samples.items():
        for index, row in enumerate(rows, start=2):
            copied = dict(row)
            copied["_sheet"] = sheet
            copied["_row"] = str(index)
            entities[company_key(copied)].append(copied)
    return dict(entities)


def first_value(rows: Iterable[dict[str, str]], *columns: str) -> str:
    for column in columns:
        for row in rows:
            value = clean(row.get(column))
            if value:
                return value
    return ""


def website_host(value: str) -> str:
    value = clean(value)
    if not value:
        return ""
    parsed = urllib.parse.urlsplit(value if "://" in value else f"https://{value}")
    return (parsed.hostname or "").lower().removeprefix("www.")


def http_url(value: str) -> str:
    value = clean(value)
    if not value:
        return ""
    candidate = value if "://" in value else f"https://{value}"
    parsed = urllib.parse.urlsplit(candidate)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return ""
    return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path or "/", parsed.query, ""))


def is_linkedin_company_url(value: str) -> bool:
    parsed = urllib.parse.urlsplit(clean(value))
    host = (parsed.hostname or "").lower()
    return (
        (host == "linkedin.com" or host.endswith(".linkedin.com"))
        and parsed.path.lower().startswith("/company/")
    )


def company_profile(key: str, rows: list[dict[str, str]]) -> dict[str, Any]:
    names: list[str] = []
    for column in (TARGET_COLUMNS["company"], "회사명(Company)", "회사명_영문(Alt Company)"):
        for row in rows:
            value = clean(row.get(column))
            if value and value not in names:
                names.append(value)
    business_no = key.removeprefix("biz:") if key.startswith("biz:") else ""
    explicit_website = first_value(rows, TARGET_COLUMNS["homepage"], "웹사이트(Website)")
    email_host = corporate_email_host(first_value(rows, "이메일(Email)"))
    return {
        "key": key,
        "rows": rows,
        "names": names,
        "company": names[0] if names else "",
        "business_no": business_no,
        "address": first_value(rows, TARGET_COLUMNS["address"], "주소(Address)"),
        "website": explicit_website or email_host,
        "website_inferred_from_email": bool(not explicit_website and email_host),
    }


def profile_needs_search(profile: dict[str, Any]) -> bool:
    return any(not clean(row.get(column)) for row in profile["rows"] for column in TARGET_COLUMNS.values())


def search_query(profile: dict[str, Any]) -> str:
    company = profile["company"]
    business_no = profile["business_no"]
    if business_no:
        return f'"{company}" "{business_no}" 기업정보 대표자 설립일 매출액'
    address_words = re.findall(r"[가-힣]{2,}(?:시|도|구|군)", profile["address"])
    location = " ".join(address_words[:2])
    domain = website_host(profile["website"])
    domain_anchor = f'"{domain}"' if domain else ""
    return clean(f'"{company}" {domain_anchor} 기업정보 대표자 설립일 매출액 {location}')


def search_payload(profile: dict[str, Any], args: argparse.Namespace) -> dict[str, Any]:
    return {
        "provider": args.provider,
        "query": search_query(profile),
        "limit": args.search_limit,
        "intent": "enterprise_db_enrichment",
    }


def crawl_payload(url: str, args: argparse.Namespace) -> dict[str, Any]:
    return {
        "url": url,
        "crawler": args.crawler,
        "respect_robots": args.respect_robots,
        "browser": "chromium",
        "allow_selenium_fallback": False,
        "base_delay_seconds": args.host_delay_seconds,
        "government_delay_seconds": max(10.0, args.host_delay_seconds),
    }


def post_json(base_url: str, path: str, payload: dict[str, Any], timeout: float) -> dict[str, Any]:
    url = f"{base_url.rstrip('/')}/{path.lstrip('/')}"
    headers = {
        "Accept": "application/json",
        "Content-Type": "application/json; charset=utf-8",
        "User-Agent": "webSearchAgent-enterprise-enricher/1.0",
    }
    token = os.environ.get("WEBSEARCH_AGENT_TOKEN", "").strip()
    if token:
        headers["Authorization"] = f"Bearer {token}"
    request = urllib.request.Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            body = response.read().decode("utf-8", errors="replace")
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")[:1500]
        raise ApiError(f"HTTP {exc.code} from {path}: {detail}") from exc
    except urllib.error.URLError as exc:
        raise ApiError(f"request failed for {path}: {exc.reason}") from exc
    try:
        data = json.loads(body)
    except json.JSONDecodeError as exc:
        raise ApiError(f"non-JSON response from {path}: {body[:300]}") from exc
    if not isinstance(data, dict):
        raise ApiError(f"non-object response from {path}")
    return data


def host_is_known(host: str) -> bool:
    return any(host == domain or host.endswith(f".{domain}") for domain in KNOWN_ENTERPRISE_DOMAINS)


def address_tokens(value: str) -> list[str]:
    tokens = re.findall(r"[가-힣A-Za-z0-9]{2,}", value)
    ignored = {"대한민국", "주식회사", "특별시", "광역시"}
    return [token for token in tokens if token not in ignored][:8]


def candidate_score(record: dict[str, Any], profile: dict[str, Any]) -> int:
    url = clean(record.get("url") or record.get("link") or record.get("canonical_url"))
    parsed = urllib.parse.urlsplit(url)
    host = (parsed.hostname or "").lower().removeprefix("www.")
    blob = clean(f"{record.get('title', '')} {record.get('snippet', '')}")
    normalized_blob = normalize_company(blob)
    score = 0
    if profile["business_no"] and profile["business_no"] in digits(blob):
        score += 100
    matched_names = [normalize_company(name) for name in profile["names"]]
    if any(len(name) >= 4 and name in normalized_blob for name in matched_names if name):
        score += 35
        if is_linkedin_company_url(url):
            score += 35
    if host_is_known(host):
        score += 25
    profile_host = website_host(profile["website"])
    if profile_host and (host == profile_host or host.endswith(f".{profile_host}")):
        score += 45
    score += min(15, sum(5 for token in address_tokens(profile["address"]) if token in blob))
    if any(marker in url.lower() for marker in IGNORED_URL_MARKERS):
        score -= 100
    return score


def ranked_candidates(response: dict[str, Any], profile: dict[str, Any]) -> list[dict[str, Any]]:
    raw_records = response.get("records")
    if not isinstance(raw_records, list):
        raw_records = response.get("items") if isinstance(response.get("items"), list) else []
    candidates: list[dict[str, Any]] = []
    seen: set[str] = set()
    for raw in raw_records:
        if not isinstance(raw, dict):
            continue
        url = http_url(clean(raw.get("url") or raw.get("link") or raw.get("canonical_url")))
        if not url or url in seen or any(marker in url.lower() for marker in IGNORED_URL_MARKERS):
            continue
        seen.add(url)
        candidate = dict(raw)
        candidate["url"] = url
        candidate["score"] = candidate_score(candidate, profile)
        candidates.append(candidate)

    homepage = http_url(profile["website"])
    if homepage and homepage not in seen:
        candidate = {"url": homepage, "title": profile["company"], "snippet": "known homepage"}
        candidate["score"] = candidate_score(candidate, profile)
        candidates.append(candidate)

    return sorted(candidates, key=lambda item: int(item["score"]), reverse=True)


def normalized_label(value: str) -> str:
    value = re.sub(r"[`*_#>\[\](){}]", "", value)
    return re.sub(r"[^0-9A-Za-z가-힣]", "", value).lower()


def canonical_field(label: str) -> str | None:
    normalized = normalized_label(label)
    if not normalized:
        return None
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            alias_normalized = normalized_label(alias)
            if normalized == alias_normalized or (
                normalized.startswith(alias_normalized) and len(normalized) <= len(alias_normalized) + 8
            ):
                return field
    return None


def sanitize_value(field: str, value: Any) -> str:
    value = html.unescape(clean(value))
    value = re.sub(r"!\[[^]]*]\([^)]*\)", "", value)
    value = re.sub(r"\[([^]]+)]\(([^)]+)\)", r"\1 \2", value)
    value = re.sub(r"<[^>]+>", " ", value)
    # Keep a leading minus on financial values; it changes the meaning.
    value = re.sub(r"^[|:：=\s]+|[|\s]+$", "", value)
    value = clean(value)
    if value.lower() in EMPTY_VALUES or len(value) > VALUE_LIMITS[field]:
        return ""
    if field == "company" and (
        value.count(",") >= 2
        or re.search(r"(?:대표자명?|사업자등록번호|본사주소|기업주소)", value)
    ):
        return ""
    if field == "business_no":
        number = digits(value)
        return number if len(number) == 10 else ""
    if field in {"founded", "revenue", "operating_profit", "net_income", "scale"} and not re.search(r"\d", value):
        return ""
    if field == "address" and len(value) < 5:
        return ""
    if field == "homepage":
        url_match = re.search(r"https?://[^\s|]+|(?:www\.)?[A-Za-z0-9.-]+\.[A-Za-z]{2,}(?:/[^\s|]*)?", value)
        if not url_match:
            return ""
        value = url_match.group(0).rstrip(".,;)")
    if field == "representative" and re.search(r"[.!?。]", value):
        return ""
    return value


def add_extracted(fields: dict[str, str], label: str, value: Any) -> None:
    field = canonical_field(label)
    if not field or field in fields:
        return
    cleaned = sanitize_value(field, value)
    if cleaned:
        fields[field] = cleaned


def extract_fields(markdown: str, text: str) -> dict[str, str]:
    fields: dict[str, str] = {}
    source = markdown or text
    for raw_line in source.splitlines():
        line = clean(raw_line)
        if not line:
            continue
        if line.count("|") >= 2:
            cells = [clean(cell) for cell in line.strip("|").split("|")]
            if len(cells) >= 2 and not all(re.fullmatch(r":?-{2,}:?", cell or "-") for cell in cells[:2]):
                add_extracted(fields, cells[0], cells[1])
        match = re.match(r"^(?:[-*#>]\s*)?(.{1,30}?)\s*[:：=]\s*(.+)$", line)
        if match:
            add_extracted(fields, match.group(1), match.group(2))
    return fields


def names_match(left: str, right: str) -> bool:
    left_normalized = normalize_company(left)
    right_normalized = normalize_company(right)
    if min(len(left_normalized), len(right_normalized)) < 4:
        return False
    return left_normalized in right_normalized or right_normalized in left_normalized


def linkedin_company_profile(crawl: dict[str, Any]) -> dict[str, Any]:
    raw_payload = crawl.get("raw_payload")
    if not isinstance(raw_payload, dict):
        return {}
    profile = raw_payload.get("profile")
    return profile if isinstance(profile, dict) else {}


def _profile_text(value: Any) -> str:
    if isinstance(value, str):
        return clean(value)
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, list):
        return ", ".join(filter(None, (_profile_text(item) for item in value[:8])))
    if isinstance(value, dict):
        preferred = ("description", "name", "city", "region", "country", "countryCode")
        texts = [_profile_text(value.get(key)) for key in preferred if value.get(key) not in (None, "")]
        return ", ".join(dict.fromkeys(filter(None, texts)))
    return ""


def extract_linkedin_company_fields(crawl: dict[str, Any]) -> dict[str, str]:
    profile = linkedin_company_profile(crawl)
    if not profile:
        return {}
    mapping = {
        "company": ("name", "companyName"),
        "homepage": ("website",),
        "address": ("locations", "headquarters"),
        "industry": ("industries", "industry"),
        "products": ("specialities", "specialties"),
        "scale": ("employeeCount", "employeeCountRange"),
        "type": ("companyType",),
        "founded": ("foundedOn", "founded"),
    }
    fields: dict[str, str] = {}
    for field, keys in mapping.items():
        for key in keys:
            value = sanitize_value(field, _profile_text(profile.get(key)))
            if value:
                fields[field] = value
                break
    return fields


def identity_match(
    profile: dict[str, Any],
    candidate: dict[str, Any],
    crawl: dict[str, Any],
    fields: dict[str, str],
) -> tuple[bool, str]:
    title = clean(crawl.get("title") or candidate.get("title"))
    page_text = clean(f"{title} {crawl.get('markdown', '')} {crawl.get('text', '')}")
    target_business_no = clean(profile.get("business_no"))
    extracted_business_no = sanitize_value("business_no", fields.get("business_no", ""))
    business_match = bool(target_business_no and target_business_no in business_numbers(page_text))
    title_name_match = any(names_match(name, title) for name in profile["names"])
    body_name_match = any(names_match(name, page_text) for name in profile["names"])
    extracted_company = sanitize_value("company", fields.get("company", ""))
    extracted_name_match = bool(extracted_company) and any(
        company_names_equivalent(name, extracted_company) for name in profile["names"]
    )

    candidate_host = website_host(clean(candidate.get("url")))
    profile_host = website_host(profile["website"])
    site_match = bool(
        profile_host and (candidate_host == profile_host or candidate_host.endswith(f".{profile_host}"))
    )
    tokens = address_tokens(profile["address"])
    address_match = sum(token in page_text for token in tokens) >= min(2, len(tokens)) if tokens else False
    trusted_directory = host_is_known(candidate_host)

    if target_business_no and extracted_business_no and target_business_no != extracted_business_no:
        return False, "business_number_conflict"
    if business_match:
        return True, "business_number"
    if is_linkedin_company_url(clean(candidate.get("url"))) and crawl.get("crawler") == "apify_linkedin":
        linkedin_profile = linkedin_company_profile(crawl)
        linkedin_name = clean(linkedin_profile.get("name") or linkedin_profile.get("companyName"))
        if not linkedin_name:
            return False, "linkedin_company_name_missing"
        linkedin_host = website_host(clean(linkedin_profile.get("website")))
        linkedin_domain_match = bool(profile_host and linkedin_host and (
            profile_host == linkedin_host
            or profile_host.endswith(f".{linkedin_host}")
            or linkedin_host.endswith(f".{profile_host}")
        ))
        if profile_host and linkedin_host and not linkedin_domain_match:
            return False, "linkedin_company_domain_conflict"
        if any(company_names_equivalent(name, linkedin_name) for name in profile["names"]):
            if linkedin_domain_match or address_match:
                return True, "linkedin_company_profile"
            return False, "linkedin_company_context_missing"
        return False, "linkedin_company_name_conflict"
    if extracted_company and not extracted_name_match:
        return False, "extracted_company_conflict"
    if site_match:
        return True, "official_website"
    if extracted_name_match and (trusted_directory or address_match or title_name_match or body_name_match):
        return True, "extracted_company_name"
    if title_name_match and (trusted_directory or address_match):
        return True, "title_company_name_with_anchor"
    if body_name_match and address_match:
        return True, "body_company_name_with_address"
    return False, "identity_not_confirmed"


def load_progress(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    records: dict[str, dict[str, Any]] = {}
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"invalid JSONL at {path}:{line_number}") from exc
            if isinstance(record, dict) and clean(record.get("key")):
                records[str(record["key"])] = record
    return records


def record_source_urls(record: dict[str, Any]) -> list[str]:
    values = list(record.get("sources") or [])
    field_sources = record.get("field_sources")
    if isinstance(field_sources, dict):
        values.extend(field_sources.values())
    return list(dict.fromkeys(clean(value) for value in values if clean(value)))


def cached_identity_anchor(profile: dict[str, Any], record: dict[str, Any]) -> str:
    urls = record_source_urls(record)
    decoded_urls = " ".join(urllib.parse.unquote(url) for url in urls)
    if profile.get("business_no") and profile["business_no"] in digits(decoded_urls):
        return "business_number_in_source_url"
    if any(names_match(name, decoded_urls) for name in profile.get("names", [])):
        return "company_name_in_source_url"
    profile_host = website_host(profile.get("website", ""))
    if profile_host and any(
        website_host(url) == profile_host or website_host(url).endswith(f".{profile_host}")
        for url in urls
    ):
        return "official_website"
    return ""


def review_cached_record(
    profile: dict[str, Any], record: dict[str, Any]
) -> tuple[dict[str, Any], bool]:
    """Recheck old cache entries created before strict company-identity validation."""
    fields = dict(record.get("fields") or {})
    known_business_no = format_business_number(profile.get("business_no"))
    added_business_no = bool(known_business_no and not fields.get("business_no"))
    if added_business_no:
        fields["business_no"] = known_business_no
    if not fields:
        return record, False

    reviewed = dict(record)
    reviewed["fields"] = fields
    reviewed["field_sources"] = dict(record.get("field_sources") or {})
    reviewed["errors"] = list(record.get("errors") or [])
    anchor = cached_identity_anchor(profile, record)
    raw_company = clean(fields.get("company"))
    extracted_company = sanitize_value("company", raw_company)

    if extracted_company and not any(
        company_names_equivalent(name, extracted_company) for name in profile.get("names", [])
    ):
        reviewed["rejected_enrichment"] = {
            "reason": "extracted_company_conflict",
            "requested_company": profile.get("company", ""),
            "extracted_company": extracted_company,
            "fields": fields,
            "sources": record_source_urls(record),
        }
        reviewed["fields"] = {"business_no": known_business_no} if known_business_no else {}
        reviewed["field_sources"] = {}
        reviewed["sources"] = []
        reviewed["status"] = "identity_conflict"
        reviewed["identity_review"] = "conflict"
        reviewed["errors"].append(
            f"cache review: requested '{profile.get('company', '')}' but extracted '{extracted_company}'"
        )
        return reviewed, True

    if raw_company and not extracted_company:
        reviewed["fields"].pop("company", None)
        reviewed["field_sources"].pop("company", None)
        reviewed["errors"].append(f"cache review: discarded malformed company field '{raw_company}'")
        if not anchor:
            reviewed["rejected_enrichment"] = {
                "reason": "company_identity_unverified",
                "requested_company": profile.get("company", ""),
                "extracted_company": raw_company,
                "fields": fields,
                "sources": record_source_urls(record),
            }
            reviewed["fields"] = {"business_no": known_business_no} if known_business_no else {}
            reviewed["field_sources"] = {}
            reviewed["sources"] = []
            reviewed["status"] = "identity_unverified"
            reviewed["identity_review"] = "unverified"
            return reviewed, True

    if extracted_company or anchor:
        reviewed["identity_review"] = "confirmed"
    else:
        reviewed["identity_review"] = "review_required"
    return reviewed, bool(added_business_no or (raw_company and not extracted_company))


def append_progress(path: Path, record: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def error_text(exc: Exception) -> str:
    return clean(str(exc))[:1200]


def throttle_host(url: str, last_request: dict[str, float], delay: float) -> None:
    host = website_host(url)
    now = time.monotonic()
    wait = max(0.0, last_request.get(host, 0.0) + delay - now)
    if wait:
        time.sleep(wait)
    last_request[host] = time.monotonic()


def process_profile(
    profile: dict[str, Any],
    args: argparse.Namespace,
    last_request: dict[str, float],
) -> dict[str, Any]:
    query = search_query(profile)
    record: dict[str, Any] = {
        "key": profile["key"],
        "company": profile["company"],
        "query": query,
        "status": "pending",
        "fields": {},
        "field_sources": {},
        "sources": [],
        "errors": [],
        "searched_at": utc_now(),
        "respect_robots": args.respect_robots,
    }
    if profile.get("business_no"):
        record["fields"]["business_no"] = format_business_number(profile["business_no"])
    if not profile_needs_search(profile):
        record["status"] = "already_complete"
        return record

    company: dict[str, Any] = {"회사명": profile["company"]}
    if len(profile["names"]) > 1:
        company["영문회사명"] = profile["names"][1]
    for key, value in (
        ("사업자등록번호", profile["business_no"]),
        ("주소", profile["address"]),
        ("홈페이지", profile["website"]),
    ):
        if value:
            company[key] = value
    try:
        response = post_json(
            args.base_url,
            "/api/enterpriseDB",
            {
                "company": company,
                "respect_robots": args.respect_robots,
                "search_limit": args.search_limit,
                "crawl_limit": args.crawl_limit,
                "linkedin_identity": True,
            },
            args.request_timeout,
        )
    except Exception as exc:
        record["status"] = "search_error"
        record["errors"].append(error_text(exc))
        return record

    items = response.get("items")
    item = items[0] if isinstance(items, list) and items and isinstance(items[0], dict) else {}
    enterprise = item.get("enterprise") if isinstance(item.get("enterprise"), dict) else {}
    enriched = set(item.get("enriched_fields") or [])
    sources = item.get("field_sources") if isinstance(item.get("field_sources"), dict) else {}
    for field, column in TARGET_COLUMNS.items():
        output_name = API_OUTPUT_NAMES.get(field, column.removeprefix("식별정보(").removesuffix(")"))
        if output_name in enriched and clean(enterprise.get(output_name)):
            record["fields"][field] = clean(enterprise[output_name])
            record["field_sources"][field] = clean(sources.get(output_name))
    record["sources"] = [
        clean(source.get("url"))
        for source in item.get("sources", [])
        if isinstance(source, dict) and clean(source.get("url"))
    ]
    raw = response.get("raw") if isinstance(response.get("raw"), dict) else {}
    record["errors"] = [
        clean(error)
        for name in ("official_errors", "search_errors", "crawl_errors")
        for error in (raw.get(name) if isinstance(raw.get(name), list) else [])
        if clean(error)
    ]
    record["query"] = clean(response.get("query")) or query
    record["status"] = clean(item.get("status") or response.get("status")) or "no_match"
    if record["sources"]:
        record["identity_review"] = "confirmed"
    return record


def metadata_for_row(row: dict[str, str], record: dict[str, Any]) -> tuple[dict[str, str], dict[str, str]]:
    updated = dict(row)
    filled: list[str] = []
    evidence: dict[str, str] = {}
    fields = record.get("fields") if isinstance(record.get("fields"), dict) else {}
    field_sources = record.get("field_sources") if isinstance(record.get("field_sources"), dict) else {}
    for field, column in TARGET_COLUMNS.items():
        value = clean(fields.get(field))
        if value and not clean(updated.get(column)):
            updated[column] = value
            filled.append(column)
            evidence[column] = clean(field_sources.get(field))

    sources = list(dict.fromkeys(value for value in evidence.values() if value))
    errors = record.get("errors") if isinstance(record.get("errors"), list) else []
    review_label = {
        "confirmed": "자동확인",
        "review_required": "수동검토 필요",
        "conflict": "회사명 충돌 - 보강 제외",
        "unverified": "회사명 미확인 - 보강 제외",
    }.get(clean(record.get("identity_review")), "")
    metadata = {
        "보강상태": "enriched" if filled else clean(record.get("status")),
        "정합성검토": review_label,
        "보강필드": ", ".join(filled),
        "보강출처": "\n".join(sources),
        "보강근거": json.dumps(evidence, ensure_ascii=False, sort_keys=True) if evidence else "",
        "보강검색어": clean(record.get("query")),
        "보강시각(UTC)": clean(record.get("searched_at")),
        "robots.txt 준수": "TRUE" if record.get("respect_robots") is True else "FALSE",
        "보강오류": " | ".join(clean(value) for value in errors if clean(value))[:3000],
    }
    return updated, metadata


def write_workbook(
    input_path: Path,
    output_path: Path,
    headers: list[str],
    samples: dict[str, list[dict[str, str]]],
    records: dict[str, dict[str, Any]],
) -> None:
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    added_fill = PatternFill("solid", fgColor="C6EFCE")
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    rejected_fill = PatternFill("solid", fgColor="F4CCCC")
    added_font = Font(bold=True, color="006100")
    review_font = Font(bold=True, color="9C6500")
    rejected_font = Font(bold=True, color="9C0006")

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        workbook = load_workbook(input_path, read_only=False, data_only=False)
    try:
        for sheet in SAMPLE_SHEETS:
            worksheet = workbook[sheet]
            output_headers = [*headers, *(column for column in TARGET_COLUMNS.values() if column not in headers)]
            added_output_headers = set(output_headers) - set(headers)
            for column_index, header in enumerate((*output_headers, *META_COLUMNS), start=1):
                cell = worksheet.cell(1, column_index, header)
                if header in added_output_headers:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="548235")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif column_index > len(output_headers):
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="305496")
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            for row_index, source_row in enumerate(samples[sheet], start=2):
                record = records[company_key(source_row)]
                updated, metadata = metadata_for_row(source_row, record)
                for column_index, header in enumerate(output_headers, start=1):
                    cell = worksheet.cell(row_index, column_index, updated.get(header, ""))
                    if (
                        header in TARGET_COLUMNS.values()
                        and not clean(source_row.get(header))
                        and clean(updated.get(header))
                    ):
                        needs_review = record.get("identity_review") == "review_required"
                        cell.fill = review_fill if needs_review else added_fill
                        cell.font = review_font if needs_review else added_font
                        field = next(name for name, column in TARGET_COLUMNS.items() if column == header)
                        source = clean((record.get("field_sources") or {}).get(field))
                        comment_lines = [
                            "WebSearch 보강값",
                            f"정합성: {metadata['정합성검토'] or '자동확인'}",
                        ]
                        if source:
                            comment_lines.append(f"출처: {source}")
                        if metadata["보강검색어"]:
                            comment_lines.append(f"검색어: {metadata['보강검색어']}")
                        cell.comment = Comment("\n".join(comment_lines), "webSearchAgent")
                for offset, header in enumerate(META_COLUMNS, start=len(output_headers) + 1):
                    cell = worksheet.cell(row_index, offset, metadata[header])
                    if record.get("identity_review") in {"conflict", "unverified"}:
                        cell.fill = rejected_fill
                        cell.font = rejected_font
                    elif record.get("identity_review") == "review_required":
                        cell.fill = review_fill
                        cell.font = review_font

            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = f"A1:{worksheet.cell(1, len(output_headers) + len(META_COLUMNS)).coordinate}"
            metadata_widths = {
                "사업자등록번호": 20,
                "정합성검토": 24,
                "보강필드": 42,
                "보강출처": 48,
                "보강근거": 48,
                "보강오류": 50,
            }
            for header in added_output_headers:
                worksheet.column_dimensions[
                    worksheet.cell(1, output_headers.index(header) + 1).column_letter
                ].width = metadata_widths.get(header, 20)
            for offset, header in enumerate(META_COLUMNS, start=len(output_headers) + 1):
                if header in metadata_widths:
                    worksheet.column_dimensions[worksheet.cell(1, offset).column_letter].width = metadata_widths[header]

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}.", suffix=".xlsx", dir=output_path.parent, delete=False
        ) as handle:
            temporary = Path(handle.name)
        try:
            workbook.save(temporary)
            os.replace(temporary, output_path)
        finally:
            temporary.unlink(missing_ok=True)
    finally:
        workbook.close()


def dry_run_report(
    args: argparse.Namespace,
    samples: dict[str, list[dict[str, str]]],
    entities: dict[str, list[dict[str, str]]],
    cached: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    profiles = [company_profile(key, rows) for key, rows in entities.items()]
    selected = profiles[: args.max_companies] if args.max_companies else profiles
    need_search = [profile for profile in selected if profile_needs_search(profile)]
    uncached = [profile for profile in need_search if args.refresh or profile["key"] not in cached]
    return {
        "mode": "dry-run",
        "network_requests_made": 0,
        "files_written": 0,
        "respect_robots": args.respect_robots,
        "input_rows": sum(len(rows) for rows in samples.values()),
        "unique_companies": len(entities),
        "selected_companies": len(selected),
        "companies_needing_search": len(need_search),
        "cached_companies_reused": len(need_search) - len(uncached),
        "cached_identity_conflicts": sum(
            record.get("identity_review") in {"conflict", "unverified"}
            for record in cached.values()
        ),
        "planned_search_requests": len(uncached),
        "maximum_planned_crawl_requests": len(uncached) * args.crawl_limit,
        "workers": args.workers,
        "input_xlsx": str(args.input_xlsx),
        "output_xlsx": str(args.output_xlsx),
        "progress_file": str(args.progress_file),
    }


def run_self_test() -> int:
    try:
        assert parse_bool("true") is True
        assert parse_bool("false") is False

        namespace = argparse.Namespace(
            crawler="auto", respect_robots=False, host_delay_seconds=2.0
        )
        payload = crawl_payload("https://example.com/company/1", namespace)
        assert payload["respect_robots"] is False

        profile = {
            "key": "name:nvidia",
            "rows": [{}],
            "names": ["NVIDIA"],
            "company": "NVIDIA",
            "business_no": "",
            "address": "",
            "website": "nvidia.com",
        }
        args = argparse.Namespace(
            base_url="https://example.com",
            respect_robots=False,
            search_limit=5,
            crawl_limit=2,
            request_timeout=10,
        )
        original_post_json = globals()["post_json"]
        globals()["post_json"] = lambda *_args, **_kwargs: {
            "query": '"NVIDIA" 기업정보',
            "status": "enriched",
            "items": [{
                "status": "enriched",
                "enterprise": {
                    "회사명": "NVIDIA",
                    "홈페이지": "https://nvidia.com",
                    "사업자번호": "123-45-67890",
                },
                "enriched_fields": ["홈페이지", "사업자번호"],
                "field_sources": {
                    "홈페이지": "https://www.linkedin.com/company/nvidia",
                    "사업자번호": "https://opendart.fss.or.kr/api/company.json?corp_code=example",
                },
                "sources": [{"url": "https://www.linkedin.com/company/nvidia"}],
            }],
            "raw": {"search_errors": [], "crawl_errors": []},
        }
        try:
            converted = process_profile(profile, args, {})
        finally:
            globals()["post_json"] = original_post_json
        assert converted["fields"] == {
            "business_no": "123-45-67890",
            "homepage": "https://nvidia.com",
        }
        assert converted["respect_robots"] is False
        reviewed, changed = review_cached_record(
            {**profile, "business_no": "1248100998"},
            {"fields": {}, "field_sources": {}, "errors": [], "sources": []},
        )
        assert changed is True
        assert reviewed["fields"]["business_no"] == "124-81-00998"

        mock = """
        | 기업명 | 테스트 주식회사 |
        | 대표자 | 홍길동 |
        | 설립일 | 2020-01-01 |
        | 매출액 | 10억원 |
        """
        fields = extract_fields(mock, mock)
        assert fields["company"] == "테스트 주식회사"
        assert fields["representative"] == "홍길동"
        assert fields["founded"] == "2020-01-01"
        assert fields["revenue"] == "10억원"

        row = {column: "" for column in TARGET_COLUMNS.values()}
        row[TARGET_COLUMNS["representative"]] = "기존대표"
        record = {
            "status": "enriched",
            "fields": {"representative": "새대표", "revenue": "10억원"},
            "field_sources": {
                "representative": "https://example.com/1",
                "revenue": "https://example.com/1",
            },
            "query": "테스트",
            "searched_at": "2026-01-01T00:00:00Z",
            "respect_robots": False,
            "errors": [],
        }
        updated, metadata = metadata_for_row(row, record)
        assert updated[TARGET_COLUMNS["representative"]] == "기존대표"
        assert updated[TARGET_COLUMNS["revenue"]] == "10억원"
        assert metadata["robots.txt 준수"] == "FALSE"

        profile = {
            "company": "nVIDIA",
            "names": ["nVIDIA"],
            "business_no": "",
            "address": "",
            "website": "nvidia.com",
        }
        matched, reason = identity_match(
            profile,
            {"url": "https://www.incruit.com/company/4126624/", "title": "리더스시스템즈"},
            {"title": "리더스시스템즈", "text": "NVIDIA 제품 공급", "markdown": ""},
            {"company": "(주)리더스시스템즈", "revenue": "100억원"},
        )
        assert matched is False and reason == "extracted_company_conflict"
        matched, reason = identity_match(
            profile,
            {"url": "https://www.nvidia.com/ko-kr/", "title": "NVIDIA"},
            {"title": "NVIDIA", "text": "NVIDIA Korea", "markdown": ""},
            {},
        )
        assert matched is True and reason == "official_website"
        assert corporate_email_host("person@example.org") == "example.org"
        assert corporate_email_host("person@gmail.com") == ""
    except (AssertionError, KeyError, ValueError) as exc:
        print(f"SELF-TEST FAILED: {exc}", file=sys.stderr)
        return 1
    print("SELF-TEST OK (no network, no files written)")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    if args.self_test:
        return run_self_test()
    validate_args(args, parser)

    try:
        headers, samples = load_inputs(args.csv_dir)
        validate_workbook(args.input_xlsx, headers, samples)
        entities = group_entities(samples)
        cached = {} if args.refresh else {
            key: record
            for key, record in load_progress(args.progress_file).items()
            if record.get("respect_robots") is args.respect_robots
            and record.get("status") != "search_error"
        }
    except (FileNotFoundError, RuntimeError, ValueError) as exc:
        parser.error(str(exc))

    profiles = [company_profile(key, rows) for key, rows in entities.items()]
    cache_corrections: list[dict[str, Any]] = []
    reviewed_cache: dict[str, dict[str, Any]] = {}
    profiles_by_key = {profile["key"]: profile for profile in profiles}
    for key, record in cached.items():
        profile = profiles_by_key.get(key)
        if not profile:
            reviewed_cache[key] = record
            continue
        reviewed, changed = review_cached_record(profile, record)
        reviewed_cache[key] = reviewed
        if changed:
            cache_corrections.append(reviewed)
    cached = reviewed_cache

    if args.dry_run:
        print(json.dumps(dry_run_report(args, samples, entities, cached), ensure_ascii=False, indent=2))
        return 0

    for correction in cache_corrections:
        append_progress(args.progress_file, correction)

    selected = profiles[: args.max_companies] if args.max_companies else profiles
    selected_keys = {profile["key"] for profile in selected}
    results = dict(cached)

    total = len(selected)
    pending: list[tuple[int, dict[str, Any]]] = []
    for index, profile in enumerate(selected, start=1):
        key = profile["key"]
        if key in results and not args.refresh:
            print(f"[{index}/{total}] cache {profile['company']}", file=sys.stderr)
            continue
        pending.append((index, profile))

    def enrich(entry: tuple[int, dict[str, Any]]) -> tuple[int, dict[str, Any], dict[str, Any]]:
        index, profile = entry
        return index, profile, process_profile(profile, args, {})

    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        for index, profile, record in executor.map(enrich, pending):
            print(f"[{index}/{total}] done {profile['company']}", file=sys.stderr)
            key = profile["key"]
            results[key] = record
            append_progress(args.progress_file, record)

    for key, rows in entities.items():
        if key in results:
            continue
        profile = company_profile(key, rows)
        status = "not_selected" if key not in selected_keys else "not_processed"
        results[key] = {
            "key": key,
            "company": profile["company"],
            "query": search_query(profile),
            "status": status,
            "fields": {},
            "field_sources": {},
            "sources": [],
            "errors": [],
            "searched_at": "",
            "respect_robots": args.respect_robots,
        }

    write_workbook(args.input_xlsx, args.output_xlsx, headers, samples, results)
    enriched_rows = sum(
        1
        for rows in samples.values()
        for row in rows
        if metadata_for_row(row, results[company_key(row)])[1]["보강필드"]
    )
    review_counts: dict[str, int] = defaultdict(int)
    for record in results.values():
        if record.get("identity_review"):
            review_counts[str(record["identity_review"])] += 1
    print(
        json.dumps(
            {
                "status": "completed",
                "output_xlsx": str(args.output_xlsx),
                "progress_file": str(args.progress_file),
                "rows_with_new_values": enriched_rows,
                "identity_review": dict(sorted(review_counts.items())),
                "cache_corrections_written": len(cache_corrections),
                "respect_robots": args.respect_robots,
                "workers": args.workers,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
